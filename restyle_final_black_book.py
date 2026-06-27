from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Engineering\Final Year Project\ManaSarathi")
DOCX = ROOT / "ManaSarathi_Final_Year_Black_Book_Report.docx"
OUTPUT = ROOT / "ManaSarathi_Final_Year_Black_Book_Report_Original_Layout.docx"
BACKUP = ROOT / "ManaSarathi_Final_Year_Black_Book_Report_before_original_layout.docx"
SPONSOR = Path(r"C:\tmp\manasarathi_report_media\01_image2.jpeg")


def delete_table(table):
    table._element.getparent().remove(table._element)


def insert_paragraph_after(paragraph, text=""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, paragraph._parent)
    if text:
        p.add_run(text)
    return p


def clear_container(container):
    for paragraph in container.paragraphs:
        paragraph.clear()


def set_run(run, size=None, bold=None, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def format_centered(paragraph, size, bold=True, color=None, before=0, after=4):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    for run in paragraph.runs:
        set_run(run, size=size, bold=bold, color=color)


def apply_original_layout(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.header_distance = Inches(0.49)
        section.footer_distance = Inches(0.49)
        clear_container(section.header)
        clear_container(section.footer)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    tokens = {
        "Heading 1": (16, "1F3864", 6, 8),
        "Heading 2": (14, "2E74B5", 11, 6),
        "Heading 3": (12.5, "2E74B5", 8, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0

    if "Chapter Label" in [s.name for s in styles]:
        style = styles["Chapter Label"]
        style.font.name = "Calibri"
        style.font.size = Pt(24)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("1F3864")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(10)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string("000000")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(5)

    for list_name in ("List Bullet", "List Number", "List Paragraph"):
        if list_name in [s.name for s in styles]:
            style = styles[list_name]
            style.font.name = "Calibri"
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing = 1.15

    if "Code Block" in [s.name for s in styles]:
        code = styles["Code Block"]
        code.font.name = "Consolas"
        code.font.size = Pt(8)
        code.paragraph_format.space_before = Pt(0)
        code.paragraph_format.space_after = Pt(0)


def restyle_front_matter(doc):
    mapping = {
        "A PROJECT REPORT ON": (11, True, None),
        "MAAN SARATHI": (18, True, "1F3864"),
        "AI-POWERED MENTAL WELLNESS PLATFORM WITH": (16, True, "1F3864"),
        "HOLISTIC HEALING AND EMOTIONAL SUPPORT": (16, True, "1F3864"),
        "SUBMITTED TO": (10.5, False, None),
        "SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE": (10.5, False, None),
        "IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF": (10.5, False, None),
        "BACHELOR OF ENGINEERING IN INFORMATION TECHNOLOGY": (14, True, None),
        "SUBMITTED BY": (11, True, None),
        "UNDER THE GUIDANCE OF": (11, True, None),
        "PROF. M. S. KALE": (11, True, None),
        "DEPARTMENT OF INFORMATION TECHNOLOGY": (11, True, None),
        "STES'S SINHGAD ACADEMY OF ENGINEERING": (11, True, None),
        "KONDHWA (BK.), PUNE - 411048": (11, True, None),
        "ACADEMIC YEAR 2025-26": (11, True, None),
        "CERTIFICATE": (15, True, "1F3864"),
        "DECLARATION": (15, True, "1F3864"),
        "ACKNOWLEDGEMENT": (15, True, "1F3864"),
        "ABSTRACT": (15, True, "1F3864"),
        "TABLE OF CONTENTS": (15, True, "1F3864"),
        "LIST OF FIGURES": (15, True, "1F3864"),
        "LIST OF TABLES": (15, True, "1F3864"),
        "ABBREVIATIONS": (15, True, "1F3864"),
    }
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if text in mapping:
            size, bold, color = mapping[text]
            format_centered(paragraph, size, bold, color, after=5)

    # Match the original cover's plain centered student-name lines.
    if doc.tables and "Student Name" in doc.tables[0].cell(0, 0).text:
        submitted = next(p for p in doc.paragraphs if p.text.strip() == "SUBMITTED BY")
        anchor = submitted
        for name in [
            "MR. ADITYA SHIRSAT     EXAM NO: ________________",
            "MR. ATHARVA LOLE       EXAM NO: ________________",
            "MISS. NEHA KAMBLE      EXAM NO: ________________",
            "MR. SOURABH SHIRKANDE  EXAM NO: ________________",
        ]:
            anchor = insert_paragraph_after(anchor, name)
            format_centered(anchor, 11, True, None, after=2)
        delete_table(doc.tables[0])


def restyle_chapter_openings(doc):
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.style.name == "Chapter Label":
            format_centered(paragraph, 24, True, "1F3864", after=4)
            # The following Heading 1 is the chapter title.
            for candidate in paragraphs[idx + 1: idx + 3]:
                if candidate.style.name == "Heading 1":
                    format_centered(candidate, 22, True, "1F3864", after=14)
                    break

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1" and paragraph.text.strip() in {
            "REFERENCES", "APPENDIX A", "APPENDIX B", "APPENDIX C"
        }:
            format_centered(paragraph, 18, True, "1F3864", after=12)


def restyle_tables(doc):
    for table in doc.tables:
        table.style = "Table Grid"
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = tc_pr.find(qn("w:shd"))
                if shd is None:
                    shd = OxmlElement("w:shd")
                    tc_pr.append(shd)
                shd.set(qn("w:fill"), "1F3864" if r_idx == 0 else "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run(run, size=9, bold=(True if r_idx == 0 else None), color=("FFFFFF" if r_idx == 0 else "000000"))


def add_sponsorship_page(doc):
    if not SPONSOR.exists() or any(p.text.strip() == "SPONSORSHIP LETTER" for p in doc.paragraphs):
        return
    declaration = next(p for p in doc.paragraphs if p.text.strip() == "DECLARATION")
    # Insert immediately before the declaration page.
    page = OxmlElement("w:p")
    declaration._p.addprevious(page)
    from docx.text.paragraph import Paragraph
    page_p = Paragraph(page, declaration._parent)
    page_p.add_run().add_break(WD_BREAK.PAGE)

    heading_xml = OxmlElement("w:p")
    declaration._p.addprevious(heading_xml)
    heading = Paragraph(heading_xml, declaration._parent)
    heading.add_run("SPONSORSHIP LETTER")
    format_centered(heading, 15, True, "1F3864", after=8)

    image_xml = OxmlElement("w:p")
    declaration._p.addprevious(image_xml)
    image_p = Paragraph(image_xml, declaration._parent)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = image_p.add_run().add_picture(str(SPONSOR), height=Inches(8.2))
    pic._inline.docPr.set("descr", "Project sponsorship letter")
    pic._inline.docPr.set("title", "Project sponsorship letter")


def main():
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)
    doc = Document(DOCX)
    apply_original_layout(doc)
    restyle_front_matter(doc)
    restyle_chapter_openings(doc)
    restyle_tables(doc)
    add_sponsorship_page(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
