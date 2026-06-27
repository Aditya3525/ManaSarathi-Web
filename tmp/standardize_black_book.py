from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor


SOURCE = Path(r"C:\Engineering\Final Year Project\ManaSarathi\tmp\ManaSarathi_working.docx")
OUTPUT = Path(r"C:\Engineering\Final Year Project\ManaSarathi\output\ManaSarathi_Black_Book_Standardized.docx")
LOGO = Path(r"C:\Engineering\Final Year Project\ManaSarathi\tmp\media_extract\word\media\image1.png")

BLUE = RGBColor(47, 84, 150)
BLACK = RGBColor(0, 0, 0)


def set_font(run, size=12, bold=None, italic=None, color=BLACK, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_para_runs(paragraph, size=12, bold=None, italic=None, color=BLACK, name="Times New Roman"):
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold, italic=italic, color=color, name=name)


def remove_all_children(el):
    for child in list(el):
        el.remove(child)


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def add_text(paragraph, text, size=12, bold=False, italic=False, color=BLACK):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def style_body_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, line=1.5):
    fmt = paragraph.paragraph_format
    paragraph.alignment = alignment
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = Inches(0.35) if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY else None
    set_para_runs(paragraph, size=12, color=BLACK)


def set_cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="000000", size="8", val="single"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def remove_table_borders(table):
    set_table_borders(table, val="nil", size="0")


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def format_front_table(table, header_fill="FFFFFF", header_font=12):
    set_table_borders(table, color="000000", size="8", val="single")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                set_para_runs(p, size=11.5, color=BLACK)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), header_fill)
        tc_pr.append(shd)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=header_font, bold=True, color=BLACK)


def add_page_border(section):
    sectPr = section._sectPr
    existing = sectPr.find(qn("w:pgBorders"))
    if existing is not None:
        sectPr.remove(existing)
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "double")
        el.set(qn("w:sz"), "18")
        el.set(qn("w:space"), "20")
        el.set(qn("w:color"), "000000")
        pg_borders.append(el)
    sectPr.append(pg_borders)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    add_page_border(section)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)

    for style_name, size, bold, align in [
        ("Heading 1", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Caption", 11, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Chapter Label", 14, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("toc 1", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("toc 2", 11, False, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            style.font.size = Pt(size)
            style.font.bold = bold
            if style.paragraph_format:
                style.paragraph_format.alignment = align


def get_paragraph_index(paragraphs, text, exact=True):
    for i, p in enumerate(paragraphs):
        content = " ".join(p.text.split())
        if exact and content == text:
            return i
        if not exact and text in content:
            return i
    raise ValueError(text)


def collect_elements_between(body, start_el, end_el=None):
    collected = []
    started = False
    for child in list(body.iterchildren()):
        if child == start_el:
            started = True
        if started:
            if end_el is not None and child == end_el:
                break
            if child.tag != qn("w:sectPr"):
                collected.append(deepcopy(child))
    return collected


def clear_body(doc):
    body = doc._element.body
    sectPr = body.sectPr
    for child in list(body):
        if child is not sectPr:
            body.remove(child)


def append_element(doc, element):
    doc._element.body.insert_element_before(deepcopy(element), "w:sectPr")


def center_paragraph(doc, text="", size=12, bold=False, color=BLACK, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def left_paragraph(doc, text="", size=12, bold=False, color=BLACK, space_before=0, space_after=8, justify=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Inches(0.35) if justify else None
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def add_logo(doc, width=Inches(2.3), after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(LOGO), width=width)
    p.paragraph_format.space_after = Pt(after)
    return p


def add_cover(doc, info):
    center_paragraph(doc, "A PROJECT REPORT ON", size=14.5, bold=True, space_before=10, space_after=12)
    center_paragraph(doc, info["name"], size=24, bold=True, space_after=8)
    center_paragraph(doc, info["subtitle1"], size=16.5, bold=True, space_after=2)
    center_paragraph(doc, info["subtitle2"], size=16.5, bold=True, space_after=12)
    add_logo(doc, width=Inches(1.95), after=8)
    center_paragraph(doc, "SUBMITTED TO THE SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE", size=12.5, space_after=4)
    center_paragraph(doc, "IN THE PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE AWARD OF THE", size=11.8, space_after=0)
    center_paragraph(doc, "DEGREE", size=11.8, space_after=6)
    center_paragraph(doc, "OF", size=11.8, space_after=6)
    center_paragraph(doc, "BACHELOR OF ENGINEERING", size=15.5, bold=True, space_after=4)
    center_paragraph(doc, "IN", size=15.5, bold=True, space_after=4)
    center_paragraph(doc, "INFORMATION TECHNOLOGY", size=15.5, bold=True, space_after=12)
    center_paragraph(doc, "SUBMITTED BY", size=13.5, bold=True, space_after=8)

    table = doc.add_table(rows=len(info["students"]), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [9.5, 7.0])
    remove_table_borders(table)
    for idx, (name, exam) in enumerate(info["students"]):
        left = table.rows[idx].cells[0]
        right = table.rows[idx].cells[1]
        set_cell_text(left, name.upper(), size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(right, f"EXAM NO: {exam}", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    center_paragraph(doc, "", space_after=6)
    center_paragraph(doc, "UNDER THE GUIDANCE OF", size=12.5, bold=True, space_after=4)
    center_paragraph(doc, info["guide"], size=12.5, bold=True, space_after=8)
    center_paragraph(doc, "DEPARTMENT OF INFORMATION TECHNOLOGY", size=14.5, bold=True, space_after=6)
    center_paragraph(doc, "STES'S SINHGAD ACADEMY OF ENGINEERING", size=12.5, bold=True, space_after=6)
    center_paragraph(doc, "KONDHWA BK, PUNE 411048", size=12.5, bold=True, space_after=6)
    center_paragraph(doc, info["year"], size=12.5, bold=True, space_after=0)
    doc.add_page_break()


def add_certificate(doc, info):
    add_logo(doc, width=Inches(2.1), after=4)
    center_paragraph(doc, "CERTIFICATE", size=20, bold=True, space_after=18)
    center_paragraph(doc, "This is to certify that the project report entitled", size=13.5, space_after=16)
    center_paragraph(doc, info["title_full"].upper(), size=15, bold=True, space_after=8)
    center_paragraph(doc, "Submitted by", size=13.5, space_after=10)

    table = doc.add_table(rows=len(info["students"]), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [9.5, 7.0])
    remove_table_borders(table)
    for idx, (name, exam) in enumerate(info["students"]):
        set_cell_text(table.rows[idx].cells[0], name.upper(), size=13, bold=True)
        set_cell_text(table.rows[idx].cells[1], f"EXAM NO: {exam}", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = left_paragraph(doc, "", size=12, space_before=10, space_after=14)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.0)
    add_text(p, "is a bonafide work has been carried out by them under the supervision of ", size=12.5)
    add_text(p, info["guide"], size=12.5, bold=True)
    add_text(p, " and it is approved for the partial fulfilment of the requirement of Savitribai Phule Pune University, for the award of the degree of ", size=12.5)
    add_text(p, "Bachelor of Engineering", size=12.5, bold=True)
    add_text(p, " (INFORMATION TECHNOLOGY).", size=12.5)

    sig = doc.add_table(rows=1, cols=3)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(sig, [6.0, 6.0, 6.0])
    remove_table_borders(sig)
    sig_labels = [
        (info["guide"], "Guide", "Department of", "Information Technology"),
        (info["hod"], "H.O.D.", "Department of", "Information Technology"),
        (info["principal"], "Principal", "SAE, Pune", ""),
    ]
    for cell, parts in zip(sig.rows[0].cells, sig_labels):
        for line_index, text in enumerate(parts):
            p = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            if text:
                add_text(p, text, size=11.5, bold=line_index < 2)

    bottom = doc.add_table(rows=1, cols=3)
    bottom.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(bottom, [6.0, 6.0, 6.0])
    remove_table_borders(bottom)

    left = bottom.rows[0].cells[0]
    for line in ("Place: Pune", "Date:      /      / 2026"):
        p = left.paragraphs[0] if not left.paragraphs[0].text else left.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, line, size=11.5)

    mid = bottom.rows[0].cells[1]
    p1 = mid.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p1, "External Examiner", size=11.5, bold=True)

    right = bottom.rows[0].cells[2]
    for idx, line in enumerate((info["guide"], "Project Coordinator")):
        p = right.paragraphs[0] if idx == 0 else right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, line, size=11.5, bold=True)

    doc.add_page_break()


def add_declaration(doc, text, members):
    center_paragraph(doc, "DECLARATION", size=20, bold=True, space_before=18, space_after=18)
    left_paragraph(doc, text, size=13, space_after=20)
    table = doc.add_table(rows=len(members) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [10.0, 6.0])
    format_front_table(table, header_fill="FFFFFF", header_font=12)
    set_cell_text(table.rows[0].cells[0], "Project Group Member", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[0].cells[1], "Signature", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, (name, _) in enumerate(members, start=1):
        set_cell_text(table.rows[idx].cells[0], name, size=12.5)
        set_cell_text(table.rows[idx].cells[1], "__________________", size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_acknowledgement(doc, paras, members):
    center_paragraph(doc, "ACKNOWLEDGEMENT", size=20, bold=True, space_before=18, space_after=18)
    for para in paras:
        if para:
            left_paragraph(doc, para, size=13, space_after=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, "Project Group Members:", size=13.5, bold=True)

    table = doc.add_table(rows=len(members), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [9.5, 7.0])
    remove_table_borders(table)
    for idx, (name, exam) in enumerate(members):
        set_cell_text(table.rows[idx].cells[0], name.upper(), size=13, bold=True)
        set_cell_text(table.rows[idx].cells[1], f"EXAM NO: {exam}", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_page_break()


def add_abstract(doc, paras):
    center_paragraph(doc, "ABSTRACT", size=20, bold=True, space_before=18, space_after=18)
    for idx, para in enumerate(paras):
        if idx < len(paras) - 1:
            left_paragraph(doc, para, size=13, space_after=10)
        else:
            p = left_paragraph(doc, "", size=13, space_before=8, space_after=0)
            p.paragraph_format.first_line_indent = Inches(0.0)
            if ":" in para:
                label, rest = para.split(":", 1)
                add_text(p, f"{label}:", size=13, bold=True, italic=True)
                add_text(p, rest, size=13, italic=True)
            else:
                add_text(p, para, size=13)
    doc.add_page_break()


def fix_frontmatter_styles(doc):
    paragraphs = doc.paragraphs
    for p in paragraphs:
        text = " ".join(p.text.split())
        if not text:
            continue
        if text in {"TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "ABBREVIATIONS"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            set_para_runs(p, size=18, bold=True, color=BLUE)
        elif p.style.name == "toc 1":
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            set_para_runs(p, size=11.5, bold=True, color=BLACK)
        elif p.style.name == "toc 2":
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(2)
            set_para_runs(p, size=11, color=BLACK)
        elif p.style.name == "Chapter Label":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            set_para_runs(p, size=14, bold=True, color=BLACK)
        elif p.style.name == "Heading 1":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            set_para_runs(p, size=16, bold=True, color=BLACK)
        elif p.style.name == "Heading 2":
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_runs(p, size=13, bold=True, color=BLACK)
        elif p.style.name == "Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            set_para_runs(p, size=11, italic=True, color=BLACK)
        elif p.style.name == "Normal":
            if text.startswith("Figure ") or text.startswith("Table "):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(4)
                set_para_runs(p, size=11.5, color=BLACK)


def fix_tables(doc):
    for idx, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    set_para_runs(p, size=11, color=BLACK)
        if idx >= 5:
            set_table_borders(table, color="000000", size="6", val="single")


def fix_footer(section):
    section.different_first_page_header_footer = True
    first_footer = section.first_page_footer
    for p in first_footer.paragraphs:
        clear_paragraph(p)
    for footer in (section.footer,):
        for p in footer.paragraphs:
            clear_paragraph(p)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))
    style_document(doc)

    paragraphs = doc.paragraphs
    tables = doc.tables

    students = []
    for row in tables[0].rows[1:]:
        name = " ".join(row.cells[0].text.split())
        exam = " ".join(row.cells[1].text.split())
        if name and exam:
            students.append((name, exam))

    info = {
        "name": "MAAN SARATHI",
        "subtitle1": "AI-POWERED MENTAL WELLNESS PLATFORM WITH",
        "subtitle2": "HOLISTIC HEALING AND EMOTIONAL SUPPORT",
        "title_full": "MaanSarathi: AI-Powered Mental Wellness Platform with Holistic Healing and Emotional Support",
        "students": students,
        "guide": "PROF. R. R. YADAV",
        "hod": "Dr. S. S. Kulkarni",
        "principal": "Dr. M. S. Rohokale",
        "year": "ACADEMIC YEAR 2025-26",
    }

    declaration_text = " ".join(paragraphs[get_paragraph_index(paragraphs, "DECLARATION") + 1].text.split())
    ack_start = get_paragraph_index(paragraphs, "ACKNOWLEDGEMENT")
    abstract_start = get_paragraph_index(paragraphs, "ABSTRACT")
    toc_start = get_paragraph_index(paragraphs, "TABLE OF CONTENTS")
    chapter_start = get_paragraph_index(paragraphs, "CHAPTER 1", exact=True)
    body_start = get_paragraph_index(paragraphs, "CHAPTER 1", exact=False)
    body_start = next(i for i, p in enumerate(paragraphs) if p.style.name == "Chapter Label" and "CHAPTER 1" in p.text)

    ack_paras = []
    idx = ack_start + 1
    while idx < abstract_start:
        text = " ".join(paragraphs[idx].text.split())
        if text:
            ack_paras.append(text)
        idx += 1

    abstract_paras = []
    idx = abstract_start + 1
    while idx < toc_start:
        text = " ".join(paragraphs[idx].text.split())
        if text:
            abstract_paras.append(text)
        idx += 1

    body = doc._element.body
    toc_block = collect_elements_between(body, paragraphs[toc_start]._element, paragraphs[body_start]._element)
    body_block = collect_elements_between(body, paragraphs[body_start]._element, None)

    clear_body(doc)

    section = doc.sections[0]
    configure_section(section)
    fix_footer(section)

    add_cover(doc, info)
    add_certificate(doc, info)
    add_declaration(doc, declaration_text, students)
    add_acknowledgement(doc, ack_paras[:4], students)
    add_abstract(doc, abstract_paras)

    for element in toc_block:
        append_element(doc, element)
    for element in body_block:
        append_element(doc, element)

    fix_frontmatter_styles(doc)
    fix_tables(doc)

    for section in doc.sections:
        configure_section(section)
        fix_footer(section)

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
