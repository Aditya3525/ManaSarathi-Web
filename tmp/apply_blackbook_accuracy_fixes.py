from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = Path(r"C:\Engineering\Final Year Project\ManaSarathi\output\ManaSarathi_Black_Book_Standardized_Corrected.docx")
OUT = Path(r"C:\Engineering\Final Year Project\ManaSarathi\output\ManaSarathi_Black_Book_Final_Accurate.docx")


def set_font(run, size=12, bold=None, italic=None, color=RGBColor(0, 0, 0), name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_text(paragraph, text, size=12, bold=False, italic=False):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)


def style_heading2(paragraph):
    paragraph.style = "Heading 2"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    set_text(paragraph, paragraph.text, size=13, bold=True)


def style_caption(paragraph):
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_font(run, size=11, italic=True)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    new_para._element = new_p
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        set_font(run)
    return new_para


def all_paragraphs(doc):
    # python-docx paragraph objects are stable enough for this one-pass update.
    return list(doc.paragraphs)


def replace_everywhere(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            for run in p.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        for run in p.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)


def main():
    doc = Document(str(SRC))

    # Official guide name.
    replace_everywhere(doc, "Prof. M. S. Kale", "Prof. R. R. Yadav")
    replace_everywhere(doc, "Prof. R.R.Yadav", "Prof. R. R. Yadav")
    replace_everywhere(doc, "PROF. R. R. YADAV", "Prof. R. R. Yadav")

    # Remove mobile as an implemented architecture claim. Browser compatibility remains valid.
    replace_everywhere(
        doc,
        "The React-based web and mobile applications provide user interaction and navigation.",
        "The React-based responsive web application provides user interaction and navigation."
    )

    # Chapter 9: restore 9.2 heading and clean sequential figure captions.
    figure_replacements = {
        "9.2 Result Summary Figure 9.4 Onboarding approach section1": "9.2 Result Summary",
        "Figure 9.5 Dashboard1": "Figure 9.6 Dashboard",
        "Figure 9.5 Assessment section": "Figure 9.7 Assessment Section",
        "Figure 9.6 Assessment Flow": "Figure 9.8 Assessment Flow",
        "Figure 9.7 Chatbot 1": "Figure 9.9 Chatbot Interface",
        "Figure 9.8 Chatbot 2": "Figure 9.10 Chatbot Response View",
        "Figure 9.8 Content Library": "Figure 9.11 Content Library",
        "Figure 9.10 Practice session": "Figure 9.12 Practice Session",
        "Figure 10.1Help Section / Therapist Booking & Overview": "Figure 9.13 Help Section / Therapist Booking and Overview",
    }
    for old, new in figure_replacements.items():
        replace_everywhere(doc, old, new)

    paragraphs = all_paragraphs(doc)
    for i, p in enumerate(paragraphs):
        text = " ".join(p.text.split())
        if text == "9.2 Result Summary":
            p.style = "Heading 2"
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            set_text(p, "9.2 Result Summary", size=13, bold=True)
            caption = insert_paragraph_after(p, "Figure 9.5 Onboarding Approach Section", "Caption")
            style_caption(caption)
            break

    for p in doc.paragraphs:
        text = " ".join(p.text.split())
        if text.startswith("Figure 9.") or text.startswith("Table 9."):
            style_caption(p)

    # Add a concise verification caveat so the testing chapter matches the current repo state.
    inserted = False
    for p in doc.paragraphs:
        if "Mental-health applications require more than technical correctness." in p.text:
            note = insert_paragraph_after(
                p,
                "Current repository verification shows broad automated backend coverage, but the latest full test run requires stabilization because two integration suites currently time out in the local environment.",
                None,
            )
            note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            note.paragraph_format.space_before = Pt(0)
            note.paragraph_format.space_after = Pt(8)
            note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            for run in note.runs:
                set_font(run, size=12)
            inserted = True
            break
    if not inserted:
        raise RuntimeError("Could not find Chapter 8 validation paragraph.")

    # Future scope should explicitly carry native mobile work as future, not implemented.
    future_scope_exists = any(
        "Native Android and iOS application" in p.text for p in doc.paragraphs
    )
    if not future_scope_exists:
        for p in doc.paragraphs:
            if "11.3 Future Scope" in p.text:
                future = insert_paragraph_after(
                    p,
                    "Native Android and iOS application with encrypted offline support, secure push notifications, and mobile-first assessment/chat workflows.",
                    "List Bullet",
                )
                for run in future.runs:
                    set_font(run, size=12)
                break
    else:
        replace_everywhere(
            doc,
            "Native Android and iOS application with encrypted offline support and notifications.",
            "Native Android and iOS application with encrypted offline support, secure push notifications, and mobile-first assessment/chat workflows.",
        )

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
