from pathlib import Path

from docx import Document


PATH = Path(r"C:\Engineering\Final Year Project\ManaSarathi\output\ManaSarathi_Black_Book_Final_Accurate.docx")


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(str(PATH))

caption_updates = {
    "Figure 9.6 Dashboard": "Figure 9.5 Dashboard",
    "Figure 9.7 Assessment Section": "Figure 9.6 Assessment Section",
    "Figure 9.8 Assessment Flow": "Figure 9.7 Assessment Flow",
    "Figure 9.9 Chatbot Interface": "Figure 9.8 Chatbot Interface",
    "Figure 9.10 Chatbot Response View": "Figure 9.9 Chatbot Response View",
    "Figure 9.11 Content Library": "Figure 9.10 Content Library",
    "Figure 9.12 Practice Session": "Figure 9.11 Practice Session",
    "Figure 9.13 Help Section / Therapist Booking and Overview": "Figure 9.12 Help Section / Therapist Booking and Overview",
}

removed = 0
renumbered = 0

for paragraph in list(doc.paragraphs):
    text = " ".join(paragraph.text.split())
    if text == "Figure 9.5 Onboarding Approach Section":
        delete_paragraph(paragraph)
        removed += 1
        continue

    if text in caption_updates:
        # Preserve the existing caption styling and replace only the visible label.
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = caption_updates[text]
        else:
            paragraph.add_run(caption_updates[text])
        renumbered += 1

if removed != 1:
    raise RuntimeError(f"Expected to remove 1 orphan onboarding caption, removed {removed}.")

if renumbered != len(caption_updates):
    raise RuntimeError(f"Expected to renumber {len(caption_updates)} captions, renumbered {renumbered}.")

doc.save(str(PATH))
print(f"Updated {PATH}")
