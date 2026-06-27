from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


PATH = Path(r"C:\Engineering\Final Year Project\ManaSarathi\output\ManaSarathi_Black_Book_Final_Accurate.docx")


def set_font(run, size=12, bold=None, italic=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph):
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)


def set_paragraph_text(paragraph, text, size=12, bold=False, italic=False):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)


def insert_paragraph_after(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    new_para._element = new_p
    new_para.style = style
    run = new_para.add_run(text)
    set_font(run, size=11, italic=True)
    new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_para.paragraph_format.space_before = Pt(6)
    new_para.paragraph_format.space_after = Pt(6)
    return new_para


def style_caption(paragraph, text):
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    set_paragraph_text(paragraph, text, size=11, italic=True)


doc = Document(str(PATH))

caption_map = {
    "Figure 9.6 Assessment Flow": "Figure 9.8 Assessment Flow",
    "Figure 9.7 Chatbot 1": "Figure 9.9 Chatbot Interface",
}

inserted_caption = False
for p in doc.paragraphs:
    text = " ".join(p.text.split())
    if text == "9.2 Result Summary Figure 9.4 Onboarding approach section1":
        p.style = "Heading 2"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        set_paragraph_text(p, "9.2 Result Summary", size=13, bold=True)
        insert_paragraph_after(p, "Figure 9.5 Onboarding Approach Section", "Caption")
        inserted_caption = True
    elif text in caption_map:
        style_caption(p, caption_map[text])

if not inserted_caption:
    raise RuntimeError("Could not split Chapter 9 result-summary heading.")

doc.save(str(PATH))
print(PATH)
